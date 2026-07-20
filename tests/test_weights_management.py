import asyncio
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path
from types import SimpleNamespace
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

from dailylogbot.config import Config
from dailylogbot.db import Database
from dailylogbot.exporters import export_pdf, export_xlsx
from dailylogbot.handlers import BotHandlers

TZ = ZoneInfo("Europe/Berlin")


@dataclass
class FakeMessage:
    text: str = ""
    replies: list[tuple[str, dict]] = field(default_factory=list)
    documents: list[dict] = field(default_factory=list)

    async def reply_text(self, text: str, **kwargs: object) -> None:
        self.replies.append((text, dict(kwargs)))

    async def reply_document(self, **kwargs: object) -> None:
        self.documents.append(dict(kwargs))


@dataclass
class FakeChat:
    async def send_action(self, action: object) -> None:
        return None


def _handlers(path: Path) -> BotHandlers:
    return BotHandlers(
        Config(
            bot_token="123456:token",
            allowed_user_id=123,
            timezone_name="Europe/Berlin",
            database_path=path,
            log_level="INFO",
        ),
        Database(path),
    )


def test_metric_weight_storage_settings_and_trash_retention(tmp_path) -> None:
    db = Database(tmp_path / "dailylog.sqlite3")
    settings = db.get_or_create_settings(123, "it")
    assert (settings.locale, settings.export_locale, settings.weight_unit) == ("it", "it", "kg")
    assert db.update_settings(123, weight_unit="lb").weight_unit == "lb"

    recorded_at = datetime(2026, 7, 8, 8, 0, tzinfo=TZ)
    weigh_in = db.add_weigh_in(123, 72_400, recorded_at)
    assert db.get_weigh_in(123, weigh_in.id).weight_grams == 72_400  # type: ignore[union-attr]
    assert db.delete_weigh_in(123, weigh_in.id, recorded_at)
    assert db.list_weigh_ins(123) == []
    assert db.list_trash(123)[0].record_type == "weigh_in"
    assert db.restore_weigh_in(123, weigh_in.id) is not None

    assert db.delete_weigh_in(123, weigh_in.id, recorded_at - timedelta(days=31))
    assert db.purge_deleted_before(recorded_at) == 1
    assert db.list_trash(123) == []


def test_weight_wizard_management_and_persisted_preferences(tmp_path) -> None:
    async def scenario() -> None:
        handlers = _handlers(tmp_path / "dailylog.sqlite3")
        context = SimpleNamespace(user_data={})
        message = FakeMessage()
        update = SimpleNamespace(
            effective_message=message,
            effective_user=SimpleNamespace(id=123, language_code="en"),
            effective_chat=FakeChat(),
        )

        for value in ("⚖️ Weigh in", "72,575"):
            message.text = value
            await handlers.message(update, context)
        saved = handlers.db.list_weigh_ins(123)[0]
        assert saved.weight_grams == 72_575
        assert "72.58 kg" in message.replies[-1][0]

        message.text = "⚙️ Settings"
        await handlers.message(update, context)
        message.text = "Pounds (lb)"
        await handlers.message(update, context)
        assert handlers.db.get_or_create_settings(123, "en").weight_unit == "lb"

        message.text = "🗂 Manage data"
        await handlers.message(update, context)
        message.text = "⚖️ Weigh-ins"
        await handlers.message(update, context)
        message.text = "📅 Today"
        await handlers.message(update, context)
        selection = message.replies[-1][1]["reply_markup"].keyboard[0][0].text
        message.text = selection
        await handlers.message(update, context)
        message.text = "🗑 Delete"
        await handlers.message(update, context)
        message.text = "🗑 Yes, delete"
        await handlers.message(update, context)
        assert handlers.db.list_weigh_ins(123) == []

        message.text = "🗂 Manage data"
        await handlers.message(update, context)
        message.text = "🗑 Trash"
        await handlers.message(update, context)
        trashed_selection = message.replies[-1][1]["reply_markup"].keyboard[0][0].text
        message.text = trashed_selection
        await handlers.message(update, context)
        message.text = "♻️ Restore"
        await handlers.message(update, context)
        assert handlers.db.list_weigh_ins(123)[0].weight_grams == 72_575

        message.text = "⚙️ Settings"
        await handlers.message(update, context)
        message.text = "Italiano"
        await handlers.message(update, context)
        restarted = _handlers(tmp_path / "dailylog.sqlite3")
        restart_context = SimpleNamespace(user_data={})
        await restarted.start(update, restart_context)
        assert "è pronto" in message.replies[-1][0]

    asyncio.run(scenario())


def test_localized_weight_only_exports(tmp_path) -> None:
    db = Database(tmp_path / "dailylog.sqlite3")
    weigh_in = db.add_weigh_in(123, 72_400, datetime(2026, 7, 8, 8, 0, tzinfo=TZ))
    output = tmp_path / "dailylog.xlsx"
    output.write_bytes(
        export_xlsx(
            [], date(2026, 7, 8), date(2026, 7, 8), TZ, weigh_ins=[weigh_in], locale="it", weight_unit="lb"
        ).getvalue()
    )
    sheet = load_workbook(output)["2026-07"]
    values = [cell.value for row in sheet.iter_rows() for cell in row if cell.value is not None]
    assert sheet["A1"].value == "REGISTRO GIORNALIERO"
    assert "PESO CORPOREO" in values
    assert any(value == "159.61 lb" for value in values)
    assert (
        export_pdf([], date(2026, 7, 8), date(2026, 7, 8), TZ, weigh_ins=[weigh_in], locale="it", weight_unit="lb")
        .getvalue()
        .startswith(b"%PDF")
    )


def test_management_paginates_record_selection(tmp_path) -> None:
    async def scenario() -> None:
        handlers = _handlers(tmp_path / "dailylog.sqlite3")
        timestamp = handlers._now().replace(hour=12, minute=0, second=0, microsecond=0)
        for offset in range(9):
            handlers.db.add_weigh_in(123, 70_000 + offset, timestamp + timedelta(minutes=offset))
        context = SimpleNamespace(user_data={})
        message = FakeMessage()
        update = SimpleNamespace(
            effective_message=message,
            effective_user=SimpleNamespace(id=123, language_code="en"),
            effective_chat=FakeChat(),
        )
        for value in ("🗂 Manage data", "⚖️ Weigh-ins", "📅 Today"):
            message.text = value
            await handlers.message(update, context)
        first_page = [button.text for row in message.replies[-1][1]["reply_markup"].keyboard for button in row]
        assert "▶" in first_page
        message.text = "▶"
        await handlers.message(update, context)
        second_page = [button.text for row in message.replies[-1][1]["reply_markup"].keyboard for button in row]
        assert "◀" in second_page
        assert any(button.startswith("#9") for button in second_page)

    asyncio.run(scenario())
