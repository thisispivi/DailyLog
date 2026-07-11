"""XLSX and PDF report generation isolated from Telegram delivery."""

from __future__ import annotations

from collections import defaultdict
from collections.abc import Iterable
from datetime import UTC, date, datetime, time, timedelta
from io import BytesIO
from xml.sax.saxutils import escape
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.cell.cell import Cell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfgen.canvas import Canvas
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from reportlab.platypus.doctemplate import BaseDocTemplate

from dailylogbot.db import Entry

NAVY = "16324F"
TEAL = "1B998B"
BLUE = "2E75B6"
AMBER = "F4A261"
VIOLET = "7B61FF"
INK = "1F2937"
MUTED = "64748B"
MIST = "F4F7FB"
LINE = "D8E1EB"
WHITE = "FFFFFF"


def export_xlsx(entries: list[Entry], start_date: date, end_date: date, tz: ZoneInfo) -> BytesIO:
    """Create a polished, month-by-month Excel log with summaries and dated sections."""
    workbook = Workbook()
    workbook.properties.creator = "DailyLog"
    workbook.properties.title = "Daily Log export"
    workbook.properties.subject = f"{start_date.isoformat()} to {end_date.isoformat()}"
    first_sheet = workbook.active
    for index, (month_key, month_start, month_end) in enumerate(_month_ranges(start_date, end_date)):
        sheet = first_sheet if index == 0 else workbook.create_sheet()
        sheet.title = month_key
        month_entries = [entry for entry in entries if month_start <= _entry_day(entry, tz) <= month_end]
        _write_month_sheet(sheet, month_entries, month_start, month_end, tz)
    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return stream


def export_pdf(entries: list[Entry], start_date: date, end_date: date, tz: ZoneInfo) -> BytesIO:
    """Create a print-ready PDF report with a clear hierarchy and coloured data table."""
    stream = BytesIO()
    document = SimpleDocTemplate(
        stream,
        pagesize=A4,
        rightMargin=10 * mm,
        leftMargin=10 * mm,
        topMargin=21 * mm,
        bottomMargin=17 * mm,
        title="Daily Log export",
        author="DailyLog",
    )
    styles = _pdf_styles()
    story: list[object] = [
        Paragraph("Daily Log", styles["title"]),
        Paragraph(f"Activity and nutrition record | {start_date:%d %b %Y} - {end_date:%d %b %Y}", styles["subtitle"]),
        Spacer(1, 6 * mm),
    ]
    for index, (month_key, month_start, month_end) in enumerate(_month_ranges(start_date, end_date)):
        month_entries = [entry for entry in entries if month_start <= _entry_day(entry, tz) <= month_end]
        story.extend(_pdf_month_section(month_key, month_entries, styles, tz))
        if index < len(_month_ranges(start_date, end_date)) - 1:
            story.append(PageBreak())
    document.build(story, onFirstPage=_draw_pdf_page, onLaterPages=_draw_pdf_page)
    stream.seek(0)
    return stream


def range_bounds(start_date: date, end_date: date, tz: ZoneInfo) -> tuple[datetime, datetime]:
    """Return timezone-aware half-open datetime bounds for an inclusive date range."""
    return (
        datetime.combine(start_date, time.min, tzinfo=tz),
        datetime.combine(end_date + timedelta(days=1), time.min, tzinfo=tz),
    )


def _write_month_sheet(
    sheet: Worksheet,
    entries: list[Entry],
    month_start: date,
    month_end: date,
    tz: ZoneInfo,
) -> None:
    """Lay out one Excel worksheet as a compact, readable monthly report."""
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.page_margins.left = 0.25
    sheet.page_margins.right = 0.25
    sheet.page_margins.top = 0.45
    sheet.page_margins.bottom = 0.45
    sheet.freeze_panes = "A9"
    widths = {"A": 16, "B": 16, "C": 24, "D": 15, "E": 15, "F": 14, "G": 56, "H": 4}
    for column_name, width in widths.items():
        sheet.column_dimensions[column_name].width = width

    sheet.merge_cells("A1:H1")
    title_cell = sheet["A1"]
    title_cell.value = "DAILY LOG"
    title_cell.font = Font(name="Aptos Display", size=22, bold=True, color=WHITE)
    title_cell.fill = PatternFill("solid", fgColor=NAVY)
    title_cell.alignment = Alignment(horizontal="left", vertical="center")
    sheet.row_dimensions[1].height = 38
    sheet.merge_cells("A2:H2")
    subtitle_cell = sheet["A2"]
    subtitle_cell.value = f"{month_start:%B %Y}  |  {month_start:%d %b} - {month_end:%d %b %Y}"
    subtitle_cell.font = Font(name="Aptos", size=11, color=MUTED)
    subtitle_cell.alignment = Alignment(horizontal="left", vertical="center")
    sheet.row_dimensions[2].height = 24
    _write_metric_cards(sheet, entries)
    headers = ["DATE", "TYPE", "ACTIVITY / MEAL", "START", "END", "DURATION", "DETAILS"]
    for column, value in enumerate(headers, start=1):
        cell = sheet.cell(row=8, column=column, value=value)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(name="Aptos", size=10, bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="left", vertical="center")
        cell.border = _bottom_border(NAVY)
    sheet.row_dimensions[8].height = 23
    row = 9
    for day, day_entries in _group_by_day(entries, tz).items():
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
        day_cell = sheet.cell(row=row, column=1, value=day.strftime("%A · %d %B %Y").upper())
        day_cell.fill = PatternFill("solid", fgColor="DCEAF7")
        day_cell.font = Font(name="Aptos", size=10, bold=True, color=BLUE)
        day_cell.alignment = Alignment(horizontal="left", vertical="center")
        day_cell.border = _bottom_border("B8D3EB")
        sheet.row_dimensions[row].height = 22
        row += 1
        for entry in _daily_order(day_entries):
            values = _xlsx_row(entry, tz)
            for column, value in enumerate(values, start=1):
                cell = sheet.cell(row=row, column=column, value=value)
                cell.font = Font(name="Aptos", size=10, color=INK)
                cell.fill = PatternFill("solid", fgColor=WHITE if row % 2 else MIST)
                cell.alignment = Alignment(
                    horizontal="left",
                    vertical="top" if column == 7 else "center",
                    wrap_text=column in {3, 7},
                )
                cell.border = _bottom_border(LINE)
            _style_category_cell(sheet.cell(row=row, column=2), entry.kind)
            sheet.row_dimensions[row].height = _xlsx_row_height(entry.text)
            row += 1
    if row == 9:
        sheet.merge_cells("A9:G9")
        empty_cell = sheet["A9"]
        empty_cell.value = "No entries recorded for this month."
        empty_cell.font = Font(name="Aptos", size=11, italic=True, color=MUTED)
        empty_cell.alignment = Alignment(horizontal="center", vertical="center")
        empty_cell.fill = PatternFill("solid", fgColor=MIST)
        sheet.row_dimensions[9].height = 30
        row = 10
    sheet.merge_cells(start_row=row + 1, start_column=1, end_row=row + 1, end_column=7)
    footer_cell = sheet.cell(row=row + 1, column=1, value="Generated by DailyLog · Private personal record")
    footer_cell.font = Font(name="Aptos", size=9, italic=True, color=MUTED)
    footer_cell.alignment = Alignment(horizontal="right", vertical="center")
    sheet.row_dimensions[row + 1].height = 22
    sheet.sheet_properties.outlinePr.summaryBelow = False
    sheet.auto_filter.ref = f"A8:G{max(9, row - 1)}"


def _write_metric_cards(sheet: Worksheet, entries: list[Entry]) -> None:
    """Write four visual metric cards above the sheet's detailed activity table."""
    metrics = [
        ("ENTRIES", len(entries), BLUE),
        ("MEALS", sum(entry.kind == "meal" for entry in entries), TEAL),
        ("TRAINING", sum(entry.kind == "training" for entry in entries), AMBER),
        ("SLEEP", sum(entry.kind == "sleep" for entry in entries), VIOLET),
    ]
    for index, (label, value, color) in enumerate(metrics):
        start_column = index * 2 + 1
        sheet.merge_cells(start_row=4, start_column=start_column, end_row=4, end_column=start_column + 1)
        sheet.merge_cells(start_row=5, start_column=start_column, end_row=5, end_column=start_column + 1)
        label_cell = sheet.cell(row=4, column=start_column, value=label)
        value_cell = sheet.cell(row=5, column=start_column, value=value)
        for cell in (label_cell, value_cell):
            cell.fill = PatternFill("solid", fgColor=color)
            cell.alignment = Alignment(horizontal="center", vertical="center")
        label_cell.font = Font(name="Aptos", size=9, bold=True, color=WHITE)
        value_cell.font = Font(name="Aptos Display", size=20, bold=True, color=WHITE)
    sheet.row_dimensions[4].height = 19
    sheet.row_dimensions[5].height = 31


def _style_category_cell(cell: Cell, kind: str) -> None:
    """Apply a distinct colour treatment for an entry category in the Excel table."""
    color = {"meal": TEAL, "training": AMBER, "sleep": VIOLET}[kind]
    cell.fill = PatternFill("solid", fgColor=color)
    cell.font = Font(name="Aptos", size=10, bold=True, color=WHITE)
    cell.alignment = Alignment(horizontal="center", vertical="center")


def _bottom_border(color: str) -> Border:
    """Create a lightweight single bottom border for clean table separation."""
    return Border(bottom=Side(style="thin", color=color))


def _xlsx_row(entry: Entry, tz: ZoneInfo) -> list[str]:
    """Return a sanitized presentation row for a single Excel entry."""
    start = entry.start_at.astimezone(tz)
    end = entry.end_at.astimezone(tz) if entry.end_at else None
    return [
        start.strftime("%d %b"),
        _kind_label(entry.kind),
        _safe_excel_text(_entry_title(entry)),
        start.strftime("%H:%M"),
        end.strftime("%H:%M") if end and entry.kind != "meal" else "",
        _duration(entry),
        _safe_excel_text(entry.text),
    ]


def _xlsx_row_height(text: str) -> float:
    """Estimate a comfortable Excel row height for wrapped user-provided details."""
    return min(90, max(24, 15 * (1 + len(text) // 85)))


def _pdf_styles() -> dict[str, ParagraphStyle]:
    """Create the compact typography system shared by the generated PDF report."""
    base = getSampleStyleSheet()
    return {
        "title": ParagraphStyle(
            "DailyLogTitle",
            parent=base["Title"],
            fontName="Helvetica-Bold",
            fontSize=25,
            leading=29,
            textColor=colors.HexColor(f"#{NAVY}"),
            spaceAfter=2,
        ),
        "subtitle": ParagraphStyle(
            "DailyLogSubtitle",
            parent=base["Normal"],
            fontName="Helvetica",
            fontSize=10,
            leading=14,
            textColor=colors.HexColor(f"#{MUTED}"),
        ),
        "month": ParagraphStyle(
            "DailyLogMonth",
            parent=base["Heading2"],
            fontName="Helvetica-Bold",
            fontSize=15,
            leading=18,
            textColor=colors.HexColor(f"#{NAVY}"),
            spaceBefore=1,
            spaceAfter=7,
        ),
        "metric_label": ParagraphStyle(
            "DailyLogMetricLabel",
            parent=base["Normal"],
            fontName="Helvetica-Bold",
            fontSize=7,
            leading=9,
            textColor=colors.white,
            alignment=TA_CENTER,
        ),
        "metric_value": ParagraphStyle(
            "DailyLogMetricValue",
            parent=base["Normal"],
            fontName="Helvetica-Bold",
            fontSize=16,
            leading=18,
            textColor=colors.white,
            alignment=TA_CENTER,
        ),
        "header": ParagraphStyle(
            "DailyLogTableHeader",
            parent=base["Normal"],
            fontName="Helvetica-Bold",
            fontSize=7.5,
            leading=9,
            textColor=colors.white,
        ),
        "day": ParagraphStyle(
            "DailyLogDay",
            parent=base["Normal"],
            fontName="Helvetica-Bold",
            fontSize=8.5,
            leading=11,
            textColor=colors.HexColor(f"#{BLUE}"),
        ),
        "cell": ParagraphStyle(
            "DailyLogCell",
            parent=base["Normal"],
            fontName="Helvetica",
            fontSize=7.6,
            leading=10,
            textColor=colors.HexColor(f"#{INK}"),
            alignment=TA_LEFT,
        ),
        "cell_bold": ParagraphStyle(
            "DailyLogCellBold",
            parent=base["Normal"],
            fontName="Helvetica-Bold",
            fontSize=7.6,
            leading=10,
            textColor=colors.HexColor(f"#{INK}"),
        ),
        "empty": ParagraphStyle(
            "DailyLogEmpty",
            parent=base["Normal"],
            fontName="Helvetica-Oblique",
            fontSize=10,
            leading=14,
            textColor=colors.HexColor(f"#{MUTED}"),
            alignment=TA_CENTER,
        ),
    }


def _pdf_month_section(
    month_key: str,
    entries: list[Entry],
    styles: dict[str, ParagraphStyle],
    tz: ZoneInfo,
) -> list[object]:
    """Build one PDF month section with summary cards and a grouped data table."""
    month_title = datetime.strptime(month_key, "%Y-%m").replace(tzinfo=UTC).strftime("%B %Y")
    section: list[object] = [
        Paragraph(month_title, styles["month"]),
        _pdf_metric_cards(entries, styles),
        Spacer(1, 5 * mm),
    ]
    if not entries:
        section.append(Paragraph("No entries recorded for this month.", styles["empty"]))
        return section
    headers = ["DATE", "TYPE", "ACTIVITY / MEAL", "TIME", "DURATION", "DETAILS"]
    data: list[list[object]] = [[Paragraph(header, styles["header"]) for header in headers]]
    commands: list[tuple[object, ...]] = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{NAVY}")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, 0), 5),
        ("BOTTOMPADDING", (0, 0), (-1, 0), 5),
        ("TOPPADDING", (0, 1), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 1), (-1, -1), 5),
        ("LINEBELOW", (0, 0), (-1, -1), 0.35, colors.HexColor(f"#{LINE}")),
    ]
    for day, day_entries in _group_by_day(entries, tz).items():
        row = len(data)
        data.append([Paragraph(day.strftime("%A | %d %B %Y").upper(), styles["day"]), "", "", "", "", ""])
        commands.extend(
            [
                ("SPAN", (0, row), (-1, row)),
                ("BACKGROUND", (0, row), (-1, row), colors.HexColor("#E8F1F9")),
                ("TOPPADDING", (0, row), (-1, row), 5),
                ("BOTTOMPADDING", (0, row), (-1, row), 4),
            ]
        )
        for entry in _daily_order(day_entries):
            row = len(data)
            data.append(_pdf_entry_row(entry, styles, tz))
            commands.append(("BACKGROUND", (1, row), (1, row), colors.HexColor(f"#{_kind_color(entry.kind)}")))
            commands.append(("TEXTCOLOR", (1, row), (1, row), colors.white))
    table = Table(data, colWidths=[55, 57, 94, 83, 56, 181], repeatRows=1, hAlign="LEFT")
    table.setStyle(TableStyle(commands))
    section.append(table)
    return section


def _pdf_metric_cards(entries: list[Entry], styles: dict[str, ParagraphStyle]) -> Table:
    """Create a four-card visual summary for a PDF month section."""
    values = [
        ("ENTRIES", len(entries), BLUE),
        ("MEALS", sum(entry.kind == "meal" for entry in entries), TEAL),
        ("TRAINING", sum(entry.kind == "training" for entry in entries), AMBER),
        ("SLEEP", sum(entry.kind == "sleep" for entry in entries), VIOLET),
    ]
    data = [
        [Paragraph(label, styles["metric_label"]) for label, _, _ in values],
        [Paragraph(str(value), styles["metric_value"]) for _, value, _ in values],
    ]
    table = Table(data, colWidths=[134.75] * 4, hAlign="LEFT")
    commands: list[tuple[object, ...]] = [
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
    ]
    for index, (_, _, color) in enumerate(values):
        commands.append(("BACKGROUND", (index, 0), (index, -1), colors.HexColor(f"#{color}")))
    table.setStyle(TableStyle(commands))
    return table


def _pdf_entry_row(entry: Entry, styles: dict[str, ParagraphStyle], tz: ZoneInfo) -> list[Paragraph]:
    """Render one entry into PDF-safe table cells without interpreting user text as markup."""
    return [
        Paragraph(_entry_day(entry, tz).strftime("%d %b"), styles["cell"]),
        Paragraph(_kind_label(entry.kind), styles["cell_bold"]),
        Paragraph(_escape_pdf(_entry_title(entry)), styles["cell_bold"]),
        Paragraph(_escape_pdf(_time_range(entry, tz)), styles["cell"]),
        Paragraph(_duration(entry), styles["cell"]),
        Paragraph(_escape_pdf(entry.text) if entry.text else "-", styles["cell"]),
    ]


def _draw_pdf_page(canvas: Canvas, document: BaseDocTemplate) -> None:
    """Draw the shared navy report header and quiet page-number footer on each page."""
    canvas.saveState()
    width, height = A4
    canvas.setFillColor(colors.HexColor(f"#{NAVY}"))
    canvas.rect(0, height - 14 * mm, width, 14 * mm, stroke=0, fill=1)
    canvas.setFillColor(colors.white)
    canvas.setFont("Helvetica-Bold", 8)
    canvas.drawString(10 * mm, height - 8.5 * mm, "DAILY LOG")
    canvas.setFillColor(colors.HexColor(f"#{MUTED}"))
    canvas.setFont("Helvetica", 7.5)
    canvas.drawRightString(width - 10 * mm, 9 * mm, f"Private personal record | Page {document.page}")
    canvas.restoreState()


def _group_by_day(entries: Iterable[Entry], tz: ZoneInfo) -> dict[date, list[Entry]]:
    """Group entries by their display day while retaining stable chronological ordering."""
    grouped: dict[date, list[Entry]] = defaultdict(list)
    for entry in entries:
        grouped[_entry_day(entry, tz)].append(entry)
    return dict(sorted(grouped.items()))


def _entry_day(entry: Entry, tz: ZoneInfo) -> date:
    """Return the day used to present an entry, placing sleep on its wake-up day."""
    timestamp = entry.end_at if entry.kind == "sleep" and entry.end_at else entry.start_at
    return timestamp.astimezone(tz).date()


def _daily_order(entries: Iterable[Entry]) -> list[Entry]:
    """Order daily records by category, conventional meal order, then clock time."""
    kind_order = {"sleep": 0, "meal": 1, "training": 2}
    meal_order = {"Breakfast": 0, "Snack 1": 1, "Lunch": 2, "Snack 2": 3, "Dinner": 4}
    return sorted(
        entries,
        key=lambda entry: (
            kind_order[entry.kind],
            meal_order.get(entry.label, 9) if entry.kind == "meal" else 0,
            entry.start_at,
            entry.id,
        ),
    )


def _entry_title(entry: Entry) -> str:
    """Return the concise primary label appropriate to an entry's category."""
    if entry.kind == "sleep":
        return "Sleep"
    return entry.label


def _kind_label(kind: str) -> str:
    """Return the human-readable English category used by portable exports."""
    return {"meal": "Meal", "training": "Training", "sleep": "Sleep"}[kind]


def _kind_color(kind: str) -> str:
    """Return the visual category colour used consistently in PDF and XLSX exports."""
    return {"meal": TEAL, "training": AMBER, "sleep": VIOLET}[kind]


def _time_range(entry: Entry, tz: ZoneInfo) -> str:
    """Format an activity's local start and end times, including overnight clarity."""
    start = entry.start_at.astimezone(tz)
    if entry.kind == "meal":
        return start.strftime("%H:%M")
    if entry.end_at is None:
        return f"{start:%H:%M} - Open"
    end = entry.end_at.astimezone(tz)
    if start.date() != end.date():
        return f"{start:%d %b %H:%M} - {end:%d %b %H:%M}"
    return f"{start:%H:%M} - {end:%H:%M}"


def _duration(entry: Entry) -> str:
    """Return a compact, non-negative elapsed duration for a closed activity."""
    if entry.kind == "meal":
        return "-"
    if entry.end_at is None:
        return "Open"
    seconds = max(0, int((entry.end_at - entry.start_at).total_seconds()))
    hours, remainder = divmod(seconds, 3600)
    minutes = remainder // 60
    if hours:
        return f"{hours}h {minutes:02d}m"
    return f"{minutes}m"


def _safe_excel_text(value: str) -> str:
    """Remove unsupported control characters and neutralize spreadsheet formula injection."""
    cleaned = "".join(character for character in value if character >= " " or character in "\n\r\t")
    return f"'{cleaned}" if cleaned.lstrip().startswith(("=", "+", "-", "@")) else cleaned


def _escape_pdf(value: str) -> str:
    """Escape text before passing it into ReportLab's lightweight paragraph markup."""
    cleaned = " ".join(value.replace("\r", " ").replace("\n", " ").split())
    return escape(cleaned)


def _month_ranges(start_date: date, end_date: date) -> list[tuple[str, date, date]]:
    """Split an inclusive range into the month-specific boundaries used by exports."""
    ranges: list[tuple[str, date, date]] = []
    current = start_date.replace(day=1)
    while current <= end_date:
        next_month = (
            current.replace(year=current.year + 1, month=1)
            if current.month == 12
            else current.replace(month=current.month + 1)
        )
        ranges.append(
            (current.strftime("%Y-%m"), max(start_date, current), min(end_date, next_month - timedelta(days=1)))
        )
        current = next_month
    return ranges
